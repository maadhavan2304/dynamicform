from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from xml.sax.saxutils import escape

from ..models import Form, FormResponse
from .validation import deserialize, display_value


def build_table(form: Form, responses: list[FormResponse]) -> tuple[list[str], list[list[str]]]:
    headers = ["Response ID", "Submitted at (UTC)", "Submitted by"] + [f.label for f in form.fields]
    rows = []
    for r in responses:
        raw = {d.field_id: d.value for d in r.details}
        rows.append(
            [r.id, r.submitted_at.strftime("%Y-%m-%d %H:%M"), r.user.name if r.user else "Anonymous"]
            + [display_value(f, deserialize(f, raw.get(f.id))) for f in form.fields]
        )
    return headers, rows


def to_xlsx(form: Form, responses: list[FormResponse]) -> bytes:
    headers, rows = build_table(form, responses)
    wb = Workbook()
    ws = wb.active
    ws.title = "Responses"
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0E7C86")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 12), 50)
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(form: Form, responses: list[FormResponse]) -> bytes:
    headers, rows = build_table(form, responses)
    styles = getSampleStyleSheet()
    cell = styles["BodyText"].clone("cell", fontSize=7, leading=9)
    head = styles["BodyText"].clone("head", fontSize=7, leading=9, textColor=colors.white, fontName="Helvetica-Bold")
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=28, bottomMargin=28,
                            title=form.title)
    data = [[Paragraph(escape(str(h)), head) for h in headers]]
    data += [[Paragraph(escape(str(c)), cell) for c in r] for r in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0E7C86")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9D1D9")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5F7")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story = [Paragraph(escape(form.title), styles["Title"]),
             Paragraph(f"{len(rows)} response(s)", styles["Normal"]), Spacer(1, 10), table]
    doc.build(story)
    return buf.getvalue()


EXPORTERS = {
    "xlsx": (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "pdf": (to_pdf, "application/pdf"),
}
